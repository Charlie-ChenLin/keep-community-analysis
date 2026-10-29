"""Inspect actual PPTX geometry and estimate changed text fit.

This is NOT an Office render or an Office-equivalent layout engine. It compares
retained slide/shape IDs with the accepted source deck, measures changed/new text
using the locally available Arial Unicode font, and reports likely clipping and
ink-box intersections. Private previews reuse *unchanged* native-chart crops
from the source's existing PDF where possible; new charts use a labelled box.
Never deliver the private preview as the presentation's PDF export.
"""
from pathlib import Path
import argparse
from collections import Counter
import hashlib
import json
import math
import re

import pymupdf as fitz
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

ROOT=Path(__file__).resolve().parents[2]
QA=Path(__file__).resolve().parent
FONT_PATH=Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
FONT=fitz.Font(fontfile=str(FONT_PATH))
EMU=914400


def pt(value):
    return float(value)*72/EMU


def font_size(paragraph):
    sizes=[r.font.size.pt for r in paragraph.runs if r.font.size is not None]
    if sizes:return max(sizes)
    if paragraph.font.size is not None:return paragraph.font.size.pt
    # A fallback is recorded, rather than claimed to resolve inherited theme text.
    return 18.0


def signature(shape):
    base=[shape.shape_type,shape.left,shape.top,shape.width,shape.height]
    if shape.has_text_frame:
        base.append(shape.text)
        base += [(r.font.name,r.font.size.pt if r.font.size is not None else None,r.font.bold)for p in shape.text_frame.paragraphs for r in p.runs]
    if shape.has_table:
        base.append([[c.text for c in row.cells]for row in shape.table.rows])
    return repr(base)


def chart_hash(shape):
    return hashlib.sha256(shape.chart.part.blob).hexdigest() if shape.has_chart else None


def char_width(char, size):
    if char=='\t':return size*2
    if not FONT.has_glyph(ord(char)) and not char.isspace():return size
    return FONT.text_length(char,fontsize=size)


def wrap_line(text,size,width,wrap):
    if not wrap:return [(text,sum(char_width(c,size)for c in text))]
    if not text:return [('',0.0)]
    tokens=re.findall(r'[A-Za-z0-9_./%+\-]+|\s+|.',text,re.S)
    lines=[];line='';used=0.0
    for token in tokens:
        tw=sum(char_width(c,size)for c in token)
        if line and used+tw>width and tw<=width:
            lines.append((line.rstrip(),used));line='';used=0.0
            token=token.lstrip();tw=sum(char_width(c,size)for c in token)
        if tw>width:
            for char in token:
                cw=char_width(char,size)
                if line and used+cw>width:
                    lines.append((line.rstrip(),used));line='';used=0.0
                line+=char;used+=cw
        else:line+=token;used+=tw
    lines.append((line.rstrip(),used))
    return lines


def text_fit(frame,bounds,unit,changed,cell=None):
    x,y,w,h=bounds
    left=pt(frame.margin_left);right=pt(frame.margin_right)
    top=pt(frame.margin_top);bottom=pt(frame.margin_bottom)
    aw=max(.1,w-left-right);ah=max(.1,h-top-bottom)
    specs=[];total=0.0;max_width=0.0;fallback=False;missing=set()
    for index,p in enumerate(frame.paragraphs):
        size=font_size(p)
        fallback |= not any(r.font.size is not None for r in p.runs) and p.font.size is None
        for char in p.text:
            if not char.isspace() and not FONT.has_glyph(ord(char)):missing.add(char)
        before=p.space_before.pt if p.space_before is not None else 0.0
        after=p.space_after.pt if p.space_after is not None and index<len(frame.paragraphs)-1 else 0.0
        sp=p.line_spacing
        if isinstance(sp,float):advance=size*sp
        elif sp is not None:advance=pt(sp)
        else:advance=size*1.15
        advance=max(size*.9,advance)
        text=p.text.replace('\v','\n')
        explicit=text.split('\n')
        lines=[]
        for s in explicit:lines.extend(wrap_line(s,size,aw,frame.word_wrap is not False))
        block_height=size*1.05+max(0,len(lines)-1)*advance
        total+=before
        specs.append({'paragraph':index,'size':size,'advance':advance,'before':before,'after':after,'lines':[a for a,_ in lines], 'widths':[b for _,b in lines],'alignment':str(p.alignment),'offset_y':total})
        total+=block_height+after
        max_width=max(max_width,max([b for _,b in lines],default=0.0))
    anchor=frame.vertical_anchor
    if anchor is None and cell is not None:anchor=cell.vertical_anchor
    dy=0.0
    if anchor==MSO_ANCHOR.MIDDLE:dy=max(0,(ah-total)/2)
    elif anchor==MSO_ANCHOR.BOTTOM:dy=max(0,ah-total)
    align=frame.paragraphs[0].alignment if frame.paragraphs else None
    dx=0.0
    if align==PP_ALIGN.CENTER:dx=max(0,(aw-max_width)/2)
    elif align==PP_ALIGN.RIGHT:dx=max(0,aw-max_width)
    excess=total-ah
    if excess>max(3.0,ah*.12):level='likely_overflow'
    elif excess>max(1.0,ah*.03):level='review_fit'
    else:level='fits_estimate'
    if max_width>aw+max(2.0,aw*.02):level='likely_overflow'
    return {'unit':unit,'changed_or_new':changed,'text':frame.text,'bounds_pt':[x,y,w,h],
            'available_width_pt':aw,'available_height_pt':ah,'estimated_text_height_pt':total,
            'estimated_max_line_width_pt':max_width,'excess_height_pt':excess,'assessment':level,
            'inherited_font_size_fallback':fallback,'unsupported_glyphs':''.join(sorted(missing)),
            'ink_bounds_pt':[x+left+dx,y+top+dy,max_width,min(total,ah)if not changed else total],
            'draw_lines':specs,'draw_origin_pt':[x+left,y+top+dy]}


def intersect(a,b):
    ax,ay,aw,ah=a;bx,by,bw,bh=b
    w=min(ax+aw,bx+bw)-max(ax,bx);h=min(ay+ah,by+bh)-max(ay,by)
    return max(0,w),max(0,h)


def draw_text(page,unit):
    ox,oy=unit['draw_origin_pt']
    aw=unit['available_width_pt']
    for p in unit['draw_lines']:
        for i,(line,width)in enumerate(zip(p['lines'],p['widths'])):
            x=ox
            if p['alignment']=='CENTER (2)':x+=max(0,(aw-width)/2)
            elif p['alignment']=='RIGHT (3)':x+=max(0,aw-width)
            y=oy+p['offset_y']+p['size']+i*p['advance']
            page.insert_text((x,y),line,fontname='audit_unicode',fontfile=str(FONT_PATH),fontsize=p['size'],color=(.12,.2,.17))


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--deck',default='Keep社区内容表达与互动反馈_老师意见修订版.pptx')
    ap.add_argument('--original',default='Keep社区内容表达与互动反馈_12分钟合并版.pptx')
    ap.add_argument('--skip-preview',action='store_true')
    args=ap.parse_args()
    target=ROOT/args.deck;source=ROOT/args.original
    if not target.exists():raise FileNotFoundError(target)
    prs=Presentation(target);old=Presentation(source)
    oldslides={s.slide_id:(i,s)for i,s in enumerate(old.slides)}
    old_chart_hashes={chart_hash(s)for slide in old.slides for s in slide.shapes if s.has_chart}
    sw,sh=pt(prs.slide_width),pt(prs.slide_height)
    source_pdf=source.with_suffix('.pdf')
    baseline_pdf=fitz.open(source_pdf)if source_pdf.exists()else None
    preview=fitz.open()
    records=[];bounds_violations=[];issues=[];overlaps=[];chart_records=[]
    changed_pages=[];unchanged_text_units=0;retained_table_units=0
    for n,slide in enumerate(prs.slides,1):
        baseline_info=oldslides.get(slide.slide_id)
        old_index,old_slide=baseline_info if baseline_info else (None,None)
        oldshapes=list(old_slide.shapes)if old_slide else []
        units=[];charts=[];pictures=[];page_changed=old_slide is None
        for shape in slide.shapes:
            x,y,w,h=map(pt,[shape.left,shape.top,shape.width,shape.height])
            if x<-.5 or y<-.5 or x+w>sw+.5 or y+h>sh+.5:
                bounds_violations.append({'page':n,'shape_id':shape.shape_id,'name':shape.name,'bounds_pt':[x,y,w,h]})
            # The accepted pptxgenjs source contains duplicate numeric shape IDs
            # on several table slides. Match content first, then type/position;
            # a shape-id dict would silently replace the earlier duplicate.
            sig=signature(shape)
            previous=next((s for s in oldshapes if signature(s)==sig),None)
            if previous is None:
                candidates=[s for s in oldshapes if s.shape_type==shape.shape_type]
                previous=min(candidates,key=lambda s:abs(s.left-shape.left)+abs(s.top-shape.top)+abs(s.width-shape.width),default=None)
            changed=previous is None or signature(shape)!=signature(previous)
            if changed:page_changed=True
            if shape.has_text_frame and shape.text.strip():
                unit=text_fit(shape.text_frame,[x,y,w,h],f'shape:{shape.shape_id}',changed)
                units.append(unit)
                if not changed:unchanged_text_units+=1
            if shape.has_table:
                cy=y
                table_changed=changed
                if not changed:retained_table_units+=1
                for ri,row in enumerate(shape.table.rows):
                    cx=x
                    for ci,cell in enumerate(row.cells):
                        cw=pt(shape.table.columns[ci].width);rh=pt(row.height)
                        # Merged cells are measured only at their merge origin.
                        if not cell.is_spanned and cell.text.strip():
                            if cell.is_merge_origin:
                                cw=sum(pt(shape.table.columns[k].width)for k in range(ci,ci+cell.span_width))
                                rh=sum(pt(shape.table.rows[k].height)for k in range(ri,ri+cell.span_height))
                            unit=text_fit(cell.text_frame,[cx,cy,cw,rh],f'table:{shape.shape_id}/r{ri+1}c{ci+1}',table_changed,cell)
                            units.append(unit)
                        cx+=pt(shape.table.columns[ci].width)
                    cy+=pt(row.height)
            if shape.has_chart:
                previous=next((s for s in oldshapes if s.has_chart and chart_hash(s)==chart_hash(shape)),previous)
                unchanged=chart_hash(shape)in old_chart_hashes
                entry={'page':n,'shape_id':shape.shape_id,'unchanged_chart_xml':unchanged,'source_page':old_index+1 if old_index is not None else None,'bounds_pt':[x,y,w,h]}
                chart_records.append(entry);charts.append((shape,previous,entry))
            if shape.shape_type==13:pictures.append(shape)
        for u in units:
            if u['changed_or_new'] and u['assessment']!='fits_estimate':issues.append({'page':n,**{k:v for k,v in u.items()if k!='draw_lines'}})
        for ai,a in enumerate(units):
            for b in units[ai+1:]:
                if not(a['changed_or_new']or b['changed_or_new']):continue
                iw,ih=intersect(a['ink_bounds_pt'],b['ink_bounds_pt'])
                if iw>2 and ih>2:
                    overlaps.append({'page':n,'a':a['unit'],'b':b['unit'],'intersection_pt':[iw,ih],'a_text':a['text'],'b_text':b['text']})
        if page_changed:changed_pages.append(n)
        records.append({'page':n,'slide_id':slide.slide_id,'source_page':old_index+1 if old_index is not None else None,'new_slide':old_slide is None,'text_units':units})
        if not args.skip_preview and page_changed:
            pg=preview.new_page(width=sw,height=sh)
            pg.draw_rect(pg.rect,color=None,fill=(1,1,1))
            # Keep the private preview visually legible on dark source slides.
            for shape in pictures:
                rect=fitz.Rect(pt(shape.left),pt(shape.top),pt(shape.left+shape.width),pt(shape.top+shape.height))
                try:pg.insert_image(rect,stream=shape.image.blob,keep_proportion=False)
                except Exception:pg.draw_rect(rect,color=(.7,.7,.7),fill=(.94,.94,.94))
            for shape,previous,entry in charts:
                rect=fitz.Rect(entry['bounds_pt'][0],entry['bounds_pt'][1],entry['bounds_pt'][0]+entry['bounds_pt'][2],entry['bounds_pt'][1]+entry['bounds_pt'][3])
                if entry['unchanged_chart_xml'] and baseline_pdf is not None and previous is not None:
                    oldrect=fitz.Rect(pt(previous.left),pt(previous.top),pt(previous.left+previous.width),pt(previous.top+previous.height))
                    pg.show_pdf_page(rect,baseline_pdf,old_index,clip=oldrect,keep_proportion=False)
                else:
                    pg.draw_rect(rect,color=(.55,.6,.57),fill=(.94,.96,.95))
                    pg.insert_text((rect.x0+5,rect.y0+14),'Native chart: geometry placeholder',fontsize=10)
            for u in units:
                if u['changed_or_new']:
                    rect=fitz.Rect(u['bounds_pt'][0],u['bounds_pt'][1],u['bounds_pt'][0]+u['bounds_pt'][2],u['bounds_pt'][1]+u['bounds_pt'][3])
                    pg.draw_rect(rect,color=(.72,.8,.75),width=.3)
                draw_text(pg,u)
            pg.insert_text((6,8),f'PRIVATE GEOMETRY PREVIEW - NOT OFFICE RENDER - revised page {n}',fontsize=5.5,color=(.6,.2,.15))
            pg.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False).save(QA/f'layout-private-{n:02d}.png')
    if len(preview):preview.save(QA/'layout-private-preview.pdf',garbage=4,deflate=True)
    stripped=[]
    for r in records:
        copied=dict(r);copied['text_units']=[{k:v for k,v in u.items()if k!='draw_lines'}for u in r['text_units']];stripped.append(copied)
    result={'deck':str(target),'deck_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'slides':len(prs.slides),
            'method':'Geometry/text-fit estimate using actual PPTX shapes, margins and explicit text sizes, measured with local Arial Unicode. Not an Office render.',
            'limitations':['Inherited theme sizes fall back to 18pt and are flagged.','Mixed font sizes use paragraph maximum size.','Line spacing and CJK wrapping approximate PowerPoint; no conclusion of exact Office fit.','Native chart labels are not re-laid out; unchanged chart XML can reuse accepted source PDF crops in private previews.','Background fills/styles are simplified in private previews; these are not a deliverable PDF.'],
            'changed_or_new_pages':changed_pages,'unchanged_text_units_reusing_baseline':unchanged_text_units,'unchanged_tables_reusing_baseline':retained_table_units,
            'native_charts':chart_records,
            'retained_unchanged_native_charts':sum(x['unchanged_chart_xml']for x in chart_records),
            'new_or_changed_native_charts':sum(not x['unchanged_chart_xml']for x in chart_records),
            'all_original_native_chart_xml_retained':old_chart_hashes.issubset({chart_hash(s)for slide in prs.slides for s in slide.shapes if s.has_chart}),
            'shape_bounds_violations':bounds_violations,'changed_text_fit_issues':issues,'changed_text_ink_overlap_candidates':overlaps,
            'page_records':stripped,'assessment':'Needs contextual review of candidates; geometry estimates alone do not certify native rendering.'}
    (QA/'slide_layout_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in result.items()if k not in ['page_records','native_charts']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
