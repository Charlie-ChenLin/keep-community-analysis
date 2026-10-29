"""Extract revised deliverables for a scoped, human-interpreted evidence check."""
from pathlib import Path
import hashlib
import json
import re
import zipfile
from lxml import etree
import pymupdf as fitz
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
RISK_PHRASES = [
    '训练和生活是两条主要内容线', '两条主要内容线仍在',
    '二者并存成为对这批有效正文较稳定的概括',
    '提示反转可能集中出现在日常问候',
    '点赞的差异更多来自“谁发”',
    '所有处理下均不显著', '已验证的新分类',
]
FORBIDDEN_PLACEHOLDERS = ['【待填写】', '待补充', 'lorem ipsum', '[insert']
FACT_ANCHORS = {
    'body_subset_size': ['5,414','5414'],
    'body_joint_p': ['0.040','.040'],
    'body_q': ['0.048','.048'],
    'vif_max': ['1.75','1.750'],
    'raw_tags': ['269'],
    'tag_recoding': ['3,342','3342'],
    'multilabel_size': ['1,088','1088'],
    'semantic_screening': ['单轮AI','AI辅助','单轮 AI','AI 辅助'],
    'provisional_codebook': ['候选','尚未验证','待验证'],
}

def norm(text):
    return re.sub(r'\s+','',text).replace('−','-').replace('，',',')

def inspect_text(text):
    compact=norm(text)
    return {'length':len(text),
            'anchors':{k:any(norm(v) in compact for v in vals)for k,vals in FACT_ANCHORS.items()},
            'risk_phrase_hits':[x for x in RISK_PHRASES if norm(x) in compact],
            'placeholder_hits':[x for x in FORBIDDEN_PLACEHOLDERS if norm(x).lower() in compact.lower()]}

def main():
    records=[]
    for path in sorted(ROOT.glob('*老师意见修订版*')):
        if path.suffix not in ['.md','.pdf','.docx','.pptx']:continue
        record={'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
        if path.suffix=='.md':
            text=path.read_text();record.update(inspect_text(text))
            (QA/(path.stem+'_md_extract.txt')).write_text(text)
        elif path.suffix=='.pdf':
            doc=fitz.open(path);text='\n'.join(p.get_text()for p in doc);record['pages']=len(doc);record.update(inspect_text(text));(QA/(path.stem+'_pdf_extract.txt')).write_text(text)
        elif path.suffix=='.docx':
            doc=Document(path);text='\n'.join(doc.element.xpath('.//w:t/text()'));record.update(inspect_text(text));(QA/(path.stem+'_docx_extract.txt')).write_text(text)
        else:
            prs=Presentation(path);slides=[]
            for n,s in enumerate(prs.slides,1):
                visible=[];chart_evidence=[]
                for shape in s.shapes:
                    if shape.has_text_frame:visible.append(shape.text)
                    if shape.has_table:visible.extend(c.text for row in shape.table.rows for c in row.cells)
                    if shape.has_chart:
                        chart=shape.chart
                        try:categories=[str(c.label)for c in chart.plots[0].categories]
                        except Exception:categories=[]
                        series=[]
                        for ss in chart.series:
                            try:values=list(ss.values)
                            except Exception:values=[]
                            series.append({'name':ss.name,'values':values})
                        chart_evidence.append({'categories':categories,'series':series})
                        visible.extend(categories)
                        for ss in series:
                            visible.append(str(ss['name']))
                            visible.extend(f'{v:.3f}'for v in ss['values']if v is not None)
                notes=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
                slides.append({'page':n,'visible':'\n'.join(visible),'notes':notes,'chart_evidence':chart_evidence,'tables':[[[c.text for c in row.cells]for row in shape.table.rows]for shape in s.shapes if shape.has_table]})
            text='\n'.join(x['visible']+'\n'+x['notes']for x in slides);record.update(inspect_text(text));record['slides']=len(slides);record['visible_only_anchors']=inspect_text('\n'.join(x['visible']for x in slides))['anchors'];(QA/'slides_extracted.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2)+'\n')
            with zipfile.ZipFile(path)as z:
                for name in z.namelist():
                    if name.endswith('.xml'):etree.fromstring(z.read(name))
                record['all_xml_well_formed']=True
                record['native_chart_parts']=len([n for n in z.namelist()if re.fullmatch(r'ppt/charts/chart\d+\.xml',n)])
        records.append(record)
    out={'scope':'Actual revised artifact text extraction and evidence anchors; flags require contextual interpretation, not proof of an error. Visual rendering is outside this script.','files':records}
    (QA/'actual_output_checks.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(out,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
