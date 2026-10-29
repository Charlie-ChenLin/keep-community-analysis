"""Independent numerical and delivery checks for the frozen reconstruction.

Uses sklearn metrics as an external implementation, never semantic adjudication.
"""
from pathlib import Path
from collections import Counter
import csv
import hashlib
import json
import math
import re
import warnings

import pymupdf as fitz
from sklearn.metrics import cohen_kappa_score,precision_score,recall_score,f1_score

ROOT=Path(__file__).resolve().parents[1]
PRIVATE=ROOT/'validation/classification_rebuild/private'


def read(path):return json.loads((ROOT/path).read_text())
def sha(path):return hashlib.sha256((ROOT/path).read_bytes()).hexdigest()
def close(a,b):return(a is None and(b is None or math.isnan(b)))or(a is not None and b is not None and math.isclose(a,b,abs_tol=1e-12))
def eligible(r,axis,code):
    return r['observation_state']=='readable'and(axis!='domains'or r['domain_state']=='specified')and code not in r['uncertain_'+axis]


def main():
    freeze=read('validation/classification_rebuild/freeze_record.json')
    assert sha('code/classification_codebook_v1.json')==freeze['codebook_sha256']
    assert sha('code/classification_recognizer_v1.json')==freeze['recognizer_sha256']
    assert all(sha(path)==digest for path,digest in freeze['artifacts'].items())
    packet=[json.loads(line)for line in(PRIVATE/'holdout_blind.jsonl').read_text().splitlines()]
    assert len(packet)==160 and all(set(row)=={'sample_id','text'}for row in packet)
    assert sha('validation/classification_rebuild/private/holdout_blind.jsonl')==freeze['holdout']['packet_sha256']
    annotations=[{r['sample_id']:r for r in read('validation/classification_rebuild/private/holdout_coder_'+c+'.json')['records']}for c in ['a','b']]
    ai=read('results/classification_ai_agreement.json')
    assert all(len(x)==160 and set(x)=={r['sample_id']for r in packet}for x in annotations)
    kappa_n=0
    for row in ai['labels']:
        code,axis=row['code'],row['axis']
        pairs=[(int(code in annotations[0][s][axis]),int(code in annotations[1][s][axis]))for s in annotations[0]if all(eligible(a[s],axis,code)for a in annotations)]
        counts=Counter(pairs)
        assert [row[x]for x in ['n11','n10','n01','n00']]==[counts[k]for k in [(1,1),(1,0),(0,1),(0,0)]]
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            expected=cohen_kappa_score([x[0]for x in pairs],[x[1]for x in pairs])if pairs else float('nan')
        assert close(row['kappa'],expected),(code,row['kappa'],expected)
        assert row['n']==len(pairs)
        denom=2*counts[(1,1)]+counts[(1,0)]+counts[(0,1)]
        assert close(row['positive_agreement'],2*counts[(1,1)]/denom if denom else None)
        kappa_n+=1
    with(ROOT/'results/rebuilt_prediction_holdout.csv').open(encoding='utf-8-sig',newline='')as f:predictions={r['sample_id']:r for r in csv.DictReader(f)}
    comparison_n=0
    for row in ai['rule_comparisons']['labels']:
        coder=annotations[0 if row['reference']=='AI_A'else 1]
        axis,code=row['axis'],row['code']
        ids=[sid for sid,r in coder.items()if eligible(r,axis,code)and predictions[sid]['pred_'+code]in ['0','1']]
        truth=[int(code in coder[s][axis])for s in ids]
        prediction=[int(predictions[s]['pred_'+code])for s in ids]
        assert row['n']==len(ids)
        assert row['TP']==sum(t==p==1 for t,p in zip(truth,prediction))
        assert row['FP']==sum(t==0 and p==1 for t,p in zip(truth,prediction))
        assert row['FN']==sum(t==1 and p==0 for t,p in zip(truth,prediction))
        if ids:
            assert close(row['precision'],precision_score(truth,prediction,zero_division=0)if sum(prediction)else None)
            assert close(row['recall'],recall_score(truth,prediction,zero_division=0)if sum(truth)else None)
            assert close(row['F1'],f1_score(truth,prediction,zero_division=0)if sum(truth)+sum(prediction)else None)
        else:assert all(row[m]is None for m in ['precision','recall','F1'])
        comparison_n+=1
    scope=read('results/classification_rebuild_scope_audit.json')
    assert all(scope['checks'].values())
    unknown=[sid for sid in annotations[0]if all(a[sid]['domain_state']=='unspecified'for a in annotations)]
    codes=[r['code']for r in ai['labels']if r['axis']=='domains']
    assigned=sum(any(predictions[s]['pred_'+c]=='1'for c in codes)for s in unknown)
    assert scope['domain_scope_diagnostics']['both_unspecified']['n']==len(unknown)==24
    assert scope['domain_scope_diagnostics']['both_unspecified']['at_least_one_domain_assigned']==assigned==18
    human=read('results/classification_human_validation_status.json')
    assert human['status']=='pending'and not human['human_agreement_evaluation_completed']
    assert human['a']['filled_cells']==human['b']['filled_cells']==0
    delivery={}
    for stem in ['分类测量重建_实证依据与验证报告','分类编码手册_v1']:
        pdf=fitz.open(ROOT/(stem+'.pdf'));text=''.join(p.get_text()for p in pdf)
        assert '\ufffd'not in text
        outside=[]
        for i,p in enumerate(pdf):
            for block in p.get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    x0,y0,x1,y1=line['bbox']
                    if x0<35 or x1>p.rect.width-32 or y0<30 or y1>p.rect.height-15:outside.append(i+1)
        assert not outside,outside
        compact=re.sub(r'\s+','',text)
        required=['AI','人工','160','80','17','15','其他']
        if stem.startswith('分类测量'):required+=['24条','18条','42条','150','41／37','153','122']
        assert all(re.sub(r'\s+','',x)in compact for x in required),[x for x in required if x not in compact]
        from docx import Document
        word=Document(ROOT/(stem+'.docx'))
        assert len(word.tables)>=8
        delivery[stem]={'pages':len(pdf),'tables_editable_docx':len(word.tables),'no_text_outside_page_margins':True,'key_claims_present':True}
        pdf.close()
    result={'status':'pass','released_packet_n':160,'frozen_definition_config_and_8_artifacts_unchanged':True,
            'blind_packet_contains_only_id_and_text':True,'AI_label_tables_checked_against_sklearn':kappa_n,
            'model_vs_AI_tables_checked_against_sklearn':comparison_n,
            'unknown_domain_assignment_recount':{'both_unspecified':len(unknown),'assigned':assigned},
            'human_validation_pending_not_synthesized':True,'delivery':delivery,
            'interpretation':'此核查验证文件、计算和导出，不验证语义正确性或人工效度。'}
    (ROOT/'validation/classification_rebuild/independent_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
