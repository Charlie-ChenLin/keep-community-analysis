"""Render the separate Phase 2 manuscript; never run analytical code.

Inputs are the Markdown manuscript and existing PNGs only. Outputs have their
own names, so the original report, slides, notebooks and numerical results stay
untouched.
"""
from pathlib import Path
import asyncio
import hashlib
import html
import json
import os
import shutil

import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from playwright.async_api import async_playwright
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[1]
STEM = '分析报告_Phase2_中文新版'
VALIDATION = ROOT / 'validation' / 'phase2_report'

CSS = """
@page { size: A4; margin: 18mm 18mm 19mm; }
* { box-sizing: border-box; }
html { background: #e9eef2; }
body { margin: 0; color: #263746; font-family: 'Microsoft YaHei',
 'PingFang SC', 'Noto Sans CJK SC', sans-serif; font-size: 10.5pt;
 line-height: 1.85; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.document { max-width: 930px; margin: 28px auto; padding: 65px 70px;
 background: white; box-shadow: 0 4px 24px #26374615; }
h1 { color: #203a50; font-size: 28pt; line-height: 1.45; margin: 10mm 0 8mm; }
h2 { color: #203a50; font-size: 20pt; line-height: 1.5; margin: 0 0 6mm;
 border-bottom: 1.5px solid #7996ad; padding-bottom: 3mm; }
h3 { color: #32617f; font-size: 13.5pt; line-height: 1.55; margin: 7mm 0 4mm; }
h1,h2,h3 { break-after: avoid; }
p { text-align: justify; margin: 0 0 4mm; orphans: 3; widows: 3; }
strong { color: #203a50; font-weight: 650; }
a { color: #32617f; text-decoration: none; overflow-wrap: anywhere; }
figure { margin: 5mm 0; break-inside: avoid; text-align: center; }
figure img { width: 100%; max-height: 123mm; object-fit: contain; }
figcaption { font-size: 9pt; line-height: 1.6; color: #5e7281; text-align: left;
 margin-top: 2mm; }
table { border-collapse: collapse; width: 100%; font-size: 9pt; line-height: 1.6;
 margin: 5mm 0; table-layout: auto; }
thead { display: table-header-group; }
tr { break-inside: avoid; }
th { text-align: left; background: #e9f0f5; font-weight: 650; }
td,th { padding: 2.2mm 2mm; border-bottom: 0.5pt solid #ccd7df; vertical-align: top; }
table p { margin: 0; }
.source { font-size: 8.7pt; color: #5e7281; line-height: 1.6; }
.cover > p { color: #5e7281; }
.cover h2 ~ p { color: #263746; }
.cover h2 { margin-top: 10mm; }
.chapter { margin-top: 15mm; }
.toc { margin: 15mm 0; }
.toc a { display: flex; justify-content: space-between; gap: 4mm;
 padding: 1mm 0; line-height: 1.55; }
.toc .page-number { flex-shrink: 0; font-variant-numeric: tabular-nums; }
.toc .level-2 { margin-left: 5mm; font-size: 9.5pt; }
.toc .level-1 { font-weight: 650; margin-top: 2mm; }
.references p { font-size: 9pt; line-height: 1.65; overflow-wrap: anywhere; }
@media print {
 html { background: white; }
 .document { max-width: none; padding: 0; margin: 0; box-shadow: none; }
 .chapter,.toc { break-before: page; margin-top: 0; }
 .toc a { padding: 0.8mm 0; }
}
"""


def read_manuscript():
    source = ROOT / f'{STEM}.md'
    soup = BeautifulSoup(markdown.markdown(source.read_text(), extensions=['tables']), 'html.parser')
    headings = []
    for index, heading in enumerate(soup.find_all(['h2', 'h3']), 1):
        heading['id'] = f'heading-{index:02d}'
        headings.append({'id': heading['id'], 'title': heading.get_text(),
                         'level': 1 if heading.name == 'h2' else 2})
    # Preserve the supplied PNGs and keep each image together with its caption.
    for paragraph in list(soup.find_all('p')):
        img = paragraph.find('img', recursive=False)
        if img is None:
            if paragraph.get_text().startswith(('来源：', '表中为', '本表为')):
                paragraph['class'] = 'source'
            continue
        path = ROOT / img['src']
        if not path.is_file():
            raise FileNotFoundError(path)
        figure = soup.new_tag('figure')
        figure.append(img.extract())
        caption = paragraph.find_next_sibling('p')
        if caption and caption.get_text().startswith('图'):
            cap = soup.new_tag('figcaption')
            cap.string = caption.get_text()
            figure.append(cap)
            caption.decompose()
        paragraph.replace_with(figure)
    return source, soup, headings


def write_html(soup, headings):
    out = BeautifulSoup('<main class="document"></main>', 'html.parser')
    main = out.main
    current = out.new_tag('section', attrs={'class': 'cover'})
    main.append(current)
    inserted_toc = False
    for node in list(soup.contents):
        if isinstance(node, NavigableString) and not node.strip():
            continue
        if getattr(node, 'name', None) == 'h2' and node.get_text() != '摘要':
            if not inserted_toc:
                nav = out.new_tag('nav', attrs={'class': 'toc'})
                title = out.new_tag('h2'); title.string = '目录'; nav.append(title)
                for heading in headings:
                    a = out.new_tag('a', href='#' + heading['id'],
                                    attrs={'class': f'level-{heading["level"]}'})
                    a.string = heading['title']; nav.append(a)
                main.append(nav); inserted_toc = True
            current = out.new_tag('section', attrs={'class': 'chapter'})
            if node.get_text() == '参考文献与证据来源':
                current['class'] = ['chapter', 'references']
            main.append(current)
        current.append(node.extract())
    title = main.h1.get_text()
    document = ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
                '<meta name="viewport" content="width=device-width, initial-scale=1">'
                f'<title>{html.escape(title)}</title><style>{CSS}</style></head>'
                f'<body>{main}</body></html>')
    (ROOT / f'{STEM}.html').write_text(document)
    return BeautifulSoup(document, 'html.parser')


def add_run_content(paragraph, node, bold=False, italic=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node)); run.bold = bold; run.italic = italic
        return
    if node.name == 'a':
        link = OxmlElement('w:hyperlink')
        href = node.get('href', '')
        if href.startswith('#'):
            link.set(qn('w:anchor'), href[1:].replace('-', '_'))
        else:
            rel = paragraph.part.relate_to(href,
                'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                is_external=True)
            link.set(qn('r:id'), rel)
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '32617F'); props.append(color)
        run.append(props); text = OxmlElement('w:t'); text.text = node.get_text()
        run.append(text); link.append(run); paragraph._p.append(link)
        return
    if node.name == 'br':
        paragraph.add_run().add_break(); return
    for child in node.children:
        add_run_content(paragraph, child, bold or node.name == 'strong', italic or node.name == 'em')


def write_docx(soup, headings):
    doc = Document(); section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9); section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(1.9); section.right_margin = Cm(1.9)
    for name, size in [('Normal', 10.5), ('Title', 26), ('Heading 1', 19),
                       ('Heading 2', 13.5), ('Caption', 9)]:
        style = doc.styles[name]; style.font.name = 'Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(size)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.6
    doc.styles['Normal'].paragraph_format.widow_control = True
    for name in ['Heading 1', 'Heading 2']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('203A50')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Keep社区内容表达与互动反馈  ·  ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    bookmark_index = 0
    for block in soup.select_one('main').children:
        if not getattr(block, 'name', None):
            continue
        if block.name == 'nav':
            doc.add_page_break(); doc.add_heading('目录', 1)
            for heading in headings:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.2
                if heading['level'] == 2: p.paragraph_format.left_indent = Cm(.5)
                a = BeautifulSoup('<a></a>', 'html.parser').a
                a['href'] = '#' + heading['id']; a.string = heading['title']
                add_run_content(p, a)
            continue
        if 'chapter' in block.get('class', []): doc.add_page_break()
        for node in block.children:
            if not getattr(node, 'name', None): continue
            if node.name in ['h1', 'h2', 'h3']:
                level = {'h1': 0, 'h2': 1, 'h3': 2}[node.name]
                p = doc.add_heading(node.get_text(), level)
                if node.has_attr('id'):
                    bookmark_index += 1
                    start = OxmlElement('w:bookmarkStart')
                    start.set(qn('w:id'), str(bookmark_index))
                    start.set(qn('w:name'), node['id'].replace('-', '_'))
                    p._p.append(start)
                    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bookmark_index))
                    p._p.append(end)
            elif node.name == 'p':
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for child in node.children: add_run_content(p, child)
                if 'source' in node.get('class', []):
                    for run in p.runs: run.font.size = Pt(9)
            elif node.name == 'figure':
                p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(ROOT / node.img['src']), width=Cm(16.8))
                doc.add_paragraph(node.figcaption.get_text(), 'Caption')
            elif node.name == 'table':
                rows = node.find_all('tr')
                table = doc.add_table(rows=0, cols=len(rows[0].find_all(['th', 'td'])))
                table.style = 'Light Shading Accent 1'
                for index, row in enumerate(rows):
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row.find_all(['th', 'td'])):
                        cell.text = value.get_text(' ', strip=True)
                        for p in cell.paragraphs:
                            p.paragraph_format.line_spacing = 1.3
                            p.paragraph_format.space_after = Pt(3)
                            for run in p.runs: run.font.size = Pt(8.5)
                    props = table.rows[-1]._tr.get_or_add_trPr()
                    no_split = OxmlElement('w:cantSplit'); props.append(no_split)
                    if index == 0:
                        props.append(OxmlElement('w:tblHeader'))
                doc.add_paragraph().paragraph_format.space_after = Pt(1)
    doc.save(ROOT / f'{STEM}.docx')


async def render_pdf(headings):
    chrome = os.environ.get('CHROME_PATH') or shutil.which('google-chrome') or shutil.which('chromium')
    mac_chrome = Path('/Applications/Google Chrome.app/Contents/MacOS/Google Chrome')
    if not chrome and mac_chrome.is_file(): chrome = str(mac_chrome)
    async with async_playwright() as pw:
        options = {'headless': True}
        if chrome: options['executable_path'] = chrome
        browser = await pw.chromium.launch(**options)
        page = await browser.new_page(viewport={'width': 1100, 'height': 1200})
        await page.goto((ROOT / f'{STEM}.html').as_uri(), wait_until='networkidle')
        await page.evaluate('document.fonts.ready')
        images = await page.locator('figure img').evaluate_all(
            '(nodes) => nodes.map(n => ({src:n.getAttribute("src"),loaded:n.complete && n.naturalWidth>0}))')
        assert all(i['loaded'] for i in images), images
        async def export():
            await page.pdf(path=str(ROOT / f'{STEM}.pdf'), print_background=True,
                           prefer_css_page_size=True, display_header_footer=True,
                           header_template='<span></span>', footer_template='''
<div style="width:100%;font-size:8px;color:#637685;margin:0 18mm;
font-family:Arial,sans-serif;display:flex;justify-content:space-between">
<span>KEEP COMMUNITY / RESEARCH REPORT</span>
<span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>''')
        await export()
        # Resolve pagination from the rendered document, then add TOC page numbers.
        # This concerns document layout only; it does not recalculate research data.
        with fitz.open(ROOT / f'{STEM}.pdf') as first:
            first_links = [link for p in first for link in p.get_links()
                           if link.get('nameddest', '').startswith('heading-')]
            first_pages = {link['nameddest']: link['page'] + 1 for link in first_links}
        document = BeautifulSoup((ROOT / f'{STEM}.html').read_text(), 'html.parser')
        for link in document.select('.toc a'):
            label = document.new_tag('span'); label.string = link.get_text(); link.clear(); link.append(label)
            number = document.new_tag('span', attrs={'class': 'page-number'})
            number.string = str(first_pages[link['href'][1:]]); link.append(number)
        (ROOT / f'{STEM}.html').write_text(str(document))
        await page.goto((ROOT / f'{STEM}.html').as_uri(), wait_until='networkidle')
        await page.evaluate('document.fonts.ready')
        await export()
        await browser.close()
    pdf = fitz.open(ROOT / f'{STEM}.pdf')
    links = [link for page in pdf for link in page.get_links() if link.get('nameddest', '').startswith('heading-')]
    by_id = {link['nameddest']: link['page'] + 1 for link in links}
    assert len(by_id) == len(headings), (len(by_id), len(headings))
    assert by_id == first_pages, 'TOC numbers changed pagination; inspect layout.'
    for heading in headings:
        target = pdf[by_id[heading['id']] - 1].get_text()
        assert ''.join(heading['title'].split()) in ''.join(target.split()), heading
    pdf.set_toc([[h['level'], h['title'], by_id[h['id']]] for h in headings])
    pdf.saveIncr()
    VALIDATION.mkdir(parents=True, exist_ok=True)
    preview_pages = {'cover': 1, 'contents': next(i + 1 for i, p in enumerate(pdf) if '目录' in p.get_text())}
    for heading in headings:
        if heading['title'].startswith(('4.4 ', '5.1 ', '5.3 ', '7.1 ')):
            preview_pages[heading['title'].split()[0].replace('.', '_')] = by_id[heading['id']]
    for label, physical in preview_pages.items():
        pix = pdf[physical - 1].get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False)
        pix.save(VALIDATION / f'{label}.png')
    return {'pages': len(pdf), 'heading_links_verified': len(headings),
            'directory_page_numbers_verified': True,
            'existing_pngs_loaded': images, 'preview_pages': preview_pages,
            'outline_entries': len(pdf.get_toc())}


def main():
    source, manuscript, headings = read_manuscript()
    html_soup = write_html(manuscript, headings)
    write_docx(html_soup, headings)
    checks = asyncio.run(render_pdf(headings))
    checks['source_markdown'] = source.name
    checks['source_sha256'] = hashlib.sha256(source.read_bytes()).hexdigest()
    checks['scope'] = 'Document rendering only; no analytical modules, raw inputs or chart generation.'
    (VALIDATION / 'render_checks.json').write_text(json.dumps(checks, ensure_ascii=False, indent=2) + '\n')
    print(f'Created separate Phase 2 report: {checks["pages"]} PDF pages, {len(checks["existing_pngs_loaded"])} existing figures.')


if __name__ == '__main__':
    main()
