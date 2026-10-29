"""Integrate the teacher-feedback evidence into a full, separately named report.

This is a delivery step. Existing analyses, manuscripts, figures and presentations
are read only. New numerical claims bind to the frozen teacher_* evidence files.
"""
from pathlib import Path
import hashlib
import html
import json
import os
import re

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / 'validation/revised_analysis_report'
STEM = '分析报告_老师意见修订版'
SOURCE = ROOT / '分析报告_Phase2_中文新版_insight_v2_修改版_排版优化.md'
os.environ.setdefault('MPLCONFIGDIR', str(QA / 'mplconfig'))

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
import markdown
import pandas as pd
import pymupdf as fitz
from bs4 import BeautifulSoup
from PIL import Image, ImageDraw
import build_phase2_report as office

LABELS = {'log_length': '正文长度 ln(1+字数)', 'self_discipline': '自律／目标词',
          'social_address': '社交／互助词', 'positive_expression': '积极表达词', 'question': '问号'}


def read(name):
    return json.loads((ROOT / 'results' / name).read_text())


def table(headers, rows):
    def cell(v): return str(v).replace('|', '／').replace('\n', ' ')
    return '\n'.join(['| ' + ' | '.join(map(cell, headers)) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |'] + ['| ' + ' | '.join(map(cell, r)) + ' |' for r in rows]) + '\n'


def f(x): return f'{float(x):.3f}'.replace('-', '−')


def ci(x): return '[' + f(x['ci_low']) + ', ' + f(x['ci_high']) + ']'


def old_section(text, start, end):
    return text[text.index(start):text.index(end)].strip()


def neutral_component_figure(topic):
    font = FontProperties(fname='/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
    terms = ['成分2：运动／时间／训练词群', '成分4：人生／生活／早安词群', '成分3：打卡／夜跑词群',
             '成分0：靠墙／膝盖词群', '成分1：公里／骑行词群']
    order = [2,4,3,0,1]
    values = [topic['original_dominant_counts'][i] / topic['documents'] * 100 for i in order]
    fig, ax = plt.subplots(figsize=(10.8,3.8))
    ax.barh(range(5), values, color='#245C4F', height=.59)
    ax.invert_yaxis()
    ax.set_yticks(range(5), terms, fontproperties=font, fontsize=12)
    for i, value in enumerate(values):
        ax.text(value+.6, i, f'{value:.1f}%', va='center', fontproperties=font, fontsize=12)
    ax.set_xlim(0,64)
    ax.set_xlabel('最高权重成分的数值份额（分母：2,470帖）', fontproperties=font, fontsize=11)
    ax.spines[['top','right']].set_visible(False)
    ax.grid(axis='x',alpha=.15); ax.set_axisbelow(True)
    fig.tight_layout()
    target = ROOT / 'figures/revised_component_shares.png'
    fig.savefig(target,dpi=180); plt.close(fig)
    return target


def manuscript():
    old = SOURCE.read_text()
    d, c = read('teacher_feedback_dictionary.json'), read('teacher_label_coverage_summary.json')
    complete, tags = read('teacher_complete_body_model.json'), read('teacher_author_tags_summary.json')
    topics, structure = read('teacher_topic_summary.json'), read('teacher_topic_structure_summary.json')
    ta = read('teacher_topic_author_support.json')
    codebook = json.loads((ROOT / 'code/revised_content_codebook.json').read_text())
    neutral_component_figure(topics)
    models = {m['model']:m for m in tags['model_fit']}
    coefs = {(x['model'],x['term']):x for x in tags['coefficients']}
    cm = complete['models']
    topic_names = list(topics['labels'].values())
    chunks = ['''# Keep社区中的字面表达与互动反馈

**老师意见修订版｜2026年10月29日**

**摘要**

这批Keep数据里，文字与点赞的关系会随比较范围改变。积极表达词命中组的原始平均点赞为57.29，未命中组为35.76；控制作者、日期和已记录情境后，估计方向转为负向。进一步区分正文是否真正可观察，五项字面特征的联合检验p从全样本的0.130，变为原始正文可见样本的0.067和清理后非空正文样本的0.040。这意味着，现有数据既不能支持固定文案稳定增赞，也不能把一种样本口径下的不显著概括为文字普遍没有关联。

内容分类需要比原先更谨慎地解释。三类词汇和问号只识别局部字面线索；5,414条非空正文中，71.9%未命中这些规则。原来提取的五个NMF成分仅覆盖2,470帖，且语义抽查和删重对照表明，成分名称、重复模板和作者产量都会影响“这批帖子主要谈什么”的判断。因此，本报告将五成分保留为探索性用词结构，不再将其份额直接解释为五类真实内容的完整分布。

评论结果呈现另一种较具体的反馈形式：4,902次明确回复中，89.0%由帖主发给其他用户；接收量相近的账号，可能分别来自53人或1,091人。评论的次数、参与人数和留言是否得到回应，需要分开理解。评论与网络两章沿用上一版数据结果，本轮修订不构成对该部分的重新全面审查。

本版把新证据整合进研究范围、变量定义、结果和结论。原M3与原五成分模型保留作历史对照；本轮补充分析均按探索性结果报告，不依据p值挑选最终模型，也不将候选语义编码规则冒充已经完成验证的新分类。
''']
    intro = old_section(old,'## 1. 研究背景、意义与研究问题','## 2. 数据来源、清洗与分析方法')
    intro = intro.replace('| RQ2 | 控制作者与帖子情境后，文字表达与帖子主题是否与点赞相关？ | 词典编码、固定效应回归、稳健性检验、TF-IDF 与NMF 主题分析 | 第4—5 章 |',
                          '| RQ2 | 在正文实际可观察的帖子中，指定字面线索与点赞有何条件关联？探索性用词成分能补充什么信息、受到哪些覆盖限制？ | 字面编码、作者和日期固定效应、正文口径及标签对照、NMF成分与分类诊断 | 第4—5章 |')
    intro = intro[:intro.index('### 1.3')] + '''### 1.3 研究预期与分析定位

目标行动用词、面向他人的称呼、积极感受用词和问号，分别提供了不同的可观察表达线索。它们值得比较，是因为读者可能在运动记录、交流邀请或日常感受中遇到这些表达；这里不把词汇命中视为作者的真实人格、情绪或行为意图。

原分析使用固定词表及全样本M3模型，并固定提取五个NMF成分。本轮针对老师提出的分类理由、共线性、非互斥、少帖作者和话题标签问题补充核查。新增的正文口径、标签控制及成分诊断都在原结果之后开展，属于探索性重分析。修订后的文字问题限定于能够观察到正文的帖子；全样本模型保留为历史对照，不因为某个结果更显著就更换原分析记录。
'''
    chunks.append(intro)
    chunks.append('''## 2. 数据来源、清洗与分析方法

### 2.1 先区分没有记录到正文、清理后为空和正文可分析

研究使用课程提供的帖子表、评论表和字段说明。9,532帖来自100位作者，帖子实际发表于2020年11—12月（UTC）；转换为北京时间后末端跨至2021年1月1日。去重后的11,595条评论最晚延续至2021年8月26日。文件名中的“2022”不代表观测年份。

| 核查项 | 发现 | 处理与解释 |
| --- | --- | --- |
| 帖子与评论主键 | 帖子ID唯一；评论有1,916行相同ID重复且无字段冲突 | 评论按ID去重，保留11,595条 |
| 原始正文 | 5,548帖有可见字符，3,984帖没有可用正文 | 未观察到正文不能视为确实没有内容 |
| 清理后正文 | 5,414帖仍非空，134帖原始可见但清理后为空 | 修订文字比较明确使用非空清理正文；同时保留其他口径对照 |
| 主题拟合与归类 | 2,478条达到拟合条件，8条贡献接近零，最终2,470帖 | 未入模帖保持未分配，不能编码为“五个主题均不存在” |
| 评论关联 | 所有去重评论均可关联到帖子，仅1,323帖采到评论 | 其余帖子的评论覆盖未知，不作全样本评论数回归 |
| 点赞与浏览 | 8,286帖点赞数高于记录浏览数 | 不计算点赞率，使用累计点赞并保留观察时长限制 |
| 时间与极端值 | 模型时间统一为北京时间，点赞有长尾 | 保留原始累计点赞，另列极端值敏感性分析 |

正文清理依次还原转义换行、移除URL、成对的课程分享后缀、方括号表情并统一空白，不把独立hashTags字段拼进正文。原全样本模型将未观察正文的文本值设0并加入正文可见标记，这只能区分记录状态和平均水平，不能恢复真实缺失内容或自动校正非随机缺失。

### 2.2 三类词汇加问号，观察的是不同表达线索

正文长度取ln(1+清理后非空白字符数)。另外四项均按字面规则编码：命中任一词或标点记1，重复出现不累加；同一帖子可以同时命中多项。它们不构成互斥内容类别。
''')
    chunks.append(table(['线索','分类理由与例子','解释边界'], [[x['label'],x['classification_axis']+'；'+ '、'.join(x['examples']),x['semantic_limit']]for x in d['definitions']]))
    chunks.append('''自律词侧重目标与行动的表述；社交词侧重称呼、感谢和交流邀请；积极词侧重表面感受与赞许；问号只是句式标点。这些选择使规则可复查，但并不证明它们是四个独立心理维度。一个“谢谢”可能兼具积极含义，“加油”既可能指自我坚持，也可能在鼓励别人；这种语义交叉需要与数值共线性分别检查。

辅助控制保留原始正文可见状态、是否有原始话题标签、表情数量、课程分享标记、媒体形式、业务类型与发帖时段。原始tag是作者选择的元数据，五个NMF成分是模型提取的词群，后续候选人工主题则需要独立语义标注，三者不能混为一类“标签”。

### 2.3 固定效应与检验沿用统一口径，比较范围单独列明

因变量为ln(1+点赞数)。原M3同时包含五项字面特征、已记录控制项、作者固定效应和北京时间发帖日期固定效应。作者固定效应扣除作者间不随帖子变化的差别；日期固定效应扣除同一天共有的差别。标准误按各模型实际包含的作者聚类，并使用有限样本修正及t／F推断。各模型的样本数和作者数分别报告，不将全样本的100位作者套到子样本。

五项核心字面系数在每个设定内用Benjamini–Hochberg方法作BH校正[6]，并作联合检验。原M3、原始正文可见模型和非空清理正文模型并列展示，避免只选择一个显著或不显著的口径。其他诊断包括正文覆盖、完整设计矩阵VIF、剔除少帖作者、标签身份控制、正文与标签拼接、极端值、作者等权、PPML及逐作者剔除。不同模型的BH家族没有合并为一次覆盖全部探索性比较的校正。

### 2.4 用词成分、主导成分与二元多标签分开处理

对至少含3个有效词的正文构建TF-IDF矩阵[4]，再以NMF提取五个用词成分[3]。原始词群名称来自关键词和正文抽查，不使用点赞。成分贡献经尺度校正并逐帖归一化，得到五项范围为0到1、合计为1的连续权重；它们不是五个0／1虚拟变量。每帖权重最大的成分只用于原描述图的互斥统计，回归仍使用连续权重。

为核对非互斥编码，另从冻结权重构建0／1多标记：某成分权重至少0.20时记1，允许同帖有多项；并用0.15与0.25作敏感性对照。阈值是本轮提出的便捷诊断，未经过语义验证。分类质量另外检查覆盖、60条分层已分配帖及20条未入模正文的单轮AI辅助语义筛查、完全重复正文去重、K=3—8和作者等权拟合目标。模型可复现或数值满秩，均不能替代标签语义验证。

评论文本、直接回复配对和有向加权网络的方法沿用上一版：根评论指向帖主，明确回复指向被回复账号，多次同向事件合并为边权。评论发生于发帖之后，不作为发帖前特征放入点赞回归。第6—7章及相关图表保留原有结果，本轮未对其进行新的完整方法复审。
''')
    third = old_section(old,'## 3. 这批数据如何影响内容与评论的比较','## 4. 文字表达与点赞的关系，取决于怎样比较')
    third = third.replace('主回归以全部帖子为比较范围，同时保留正文可用状态；主题分析则聚焦满足文本条件的正文。两者分别描述全样本中记录到的文字线索和有效正文中的内容结构。图片文字、视频语音属于正文记录以外的表达形式，文本覆盖的具体范围见附录A。',
        '原M3以全部帖子为历史比较范围，修订后的字面分析同时列出原始正文可见与清理后非空的样本。NMF最终只有2,470帖获得可用权重，占全样本25.9%、非空正文45.6%；另有2,944条非空正文未进入该主题表。未入模不等于没有主题，正文缺失也不等于图片和视频没有内容。图片文字和视频语音尚未观察，不能用现有文字类目补齐。\n\n![补充图S1](figures/teacher_text_coverage.png)\n\n图S1　全样本按正文与原主题模型的覆盖状态划分，三类互斥，合计9,532帖。它显示分析能看到哪些记录，不是五类之外的语义比例。')
    chunks.append(third)
    chunks.append('''## 4. 字面表达的关联对正文可观测范围敏感

### 4.1 原始均值有差异，但词典只能覆盖部分表达

以下原始均值沿用历史9,532帖的编码比较，“未命中”组包含没有观察到正文的记录，不能视为已经确认这些帖子不使用相应用词。未调整时，四项字面线索的命中组平均点赞均较高：自律／目标词63.67对34.36，社交／互助词89.96对36.16，积极表达词57.29对35.76，问号111.15对36.07。两组的作者、媒体和发帖情境并不相同，这些均值描述的是组别构成与反馈的共同差别，不能直接解释为措辞效果。

![图4](figures/03_text_features.png)

图4　指定字面特征的命中数与高频词。它们可重叠，词频以含该词的帖子数计，每帖至多一次。
''')
    chunks.append(table(['字面线索','命中帖数','全部帖子占比','有作者内变化的作者数','完整原M3 VIF'], [[x['label'],x['n_hit'],f"{x['share_all']:.1%}",x['authors_with_within_variation'],f(x['vif_full_main_model'])]for x in d['definitions']]))
    chunks.append('''非空清理正文中，只有1,523帖命中至少一项，3,891帖没有命中任何规则，占71.9%。全样本未命中数为8,009，其中包含未观察到正文的记录。这不是“多数帖子没有表达”，而是说明指定词表只捕捉了局部线索。RQ2的结论因而必须限定为这些规则的关联，不能概括全部内容的作用。

四项全样本最大两两φ相关为0.119，作者内最大相关为0.085；完整M3的四项VIF为1.11—1.75，设计矩阵9,532×182且满秩。182帖同时命中至少两项。现有数值未呈现明显共线性，不支持仅因担心重复而合并规则；低VIF也不能证明语义独立或量表有效。[10]

### 4.2 原M3的估计可复算，但不能替代可观测正文中的比较

原全样本M3中，积极表达系数为−0.086，原始95%区间为[−0.155, −0.016]，原始p=0.016；对五项同时作BH校正后q=0.079。五项联合p=0.130，新增R²为0.065个百分点。这些是原设定下可以复算的结果。

![图5](figures/04_main_coefficients.png)

图5　原全样本M3的系数与原始95%作者聚类区间，N=9,532、100位作者；非空正文结果另列下表。

原始算术均值与模型的对数条件关联也不在同一尺度，不能将方向差别单独归因于作者固定效应或某一个控制项。把研究对象限定为正文实际可见的帖子后，比较范围发生了变化。原始正文可见的5,548帖仍包含134条清理后为空的记录；非空清理正文为5,414帖。三个口径的结果并列如下：
''')
    chunks.append(table(['正文口径','帖数／作者数','五项联合p','长度BH q','积极表达BH q'], [[name,f"{m['n']:,}／{m['authors']}",f(m['joint_five_p']),f(next(x['q_bh']for x in m['terms']if x['term']=='log_length')),f(next(x['q_bh']for x in m['terms']if x['term']=='positive_expression'))]for name,m in zip(['原全样本＋缺失标记','原始正文可见','清理后正文非空'],cm)]))
    chunks.append('''在非空正文模型中，长度系数为0.046、积极表达系数为−0.089，两项在该模型内的BH q均约0.048，整组p约0.040。不能因此挑选这个模型来宣布稳定效果，也不能继续把全样本p=0.130概括为所有正文处理都没有关联。该重分析是在原结果之后开展，且没有对所有探索性模型统一校正；结果靠近阈值，仍受到正文可观察范围的选择限制。

修订后的问题明确聚焦“可观察的正文中的指定线索”。选择这个范围的理由是变量确实能够被测量，不是它的p值更小。原全样本模型作为历史对照保留。两者都缺少实际曝光、完整媒体内容和统一累计观察时长，不能识别词汇的因果作用。

### 4.3 少帖作者没有主导原字面结果，稀少特征的支持仍有限

全样本100位作者的发帖数最少1、中位数66、最多783。2位作者只发1帖；不超过10帖的14位作者合计80帖，占0.84%。单帖作者没有作者内变化，不为共同斜率提供这种信息，但少帖存在本身不使固定效应整体失效。

分别剔除发帖少于5帖或10帖的作者，原字面模型得到以下结果：
''')
    chunks.append(table(['保留规则','帖数／作者数','五项联合p','积极表达β／BH q'], [[name,f"{models[key]['n']:,}／{models[key]['authors']}",f(models[key]['joint_five_p']),f(coefs[key,'positive_expression']['coef'])+'／'+f(coefs[key,'positive_expression']['q_bh_five'])]for key,name in [('baseline_body','原全样本'),('authors_ge5','每位作者至少5帖'),('authors_ge10','每位作者至少10帖')]]))
    chunks.append('''这些门槛没有实质改变原M3的统计判断，提示少帖作者不是这项结果的主要来源。不过，长度在93位作者内有变化，社交词只有42位、问号只有32位；不能把100个作者理解为每个特征都有100个充分的变化来源。固定效应也不能排除同一作者随帖变化的曝光和内容差别。

### 4.4 原始话题标签值得控制，但拼进正文会改变测量

5,209帖有原始tag，占54.6%，共269种；其中1,517帖没有原始正文。原M3已经控制“有无标签”，并非没有考虑tag，但没有控制具体标签身份。

本轮只依据频数与作者覆盖，选择至少30帖、至少5位作者使用的9个tag身份作为补充控制，并保留原正文编码。另把正文与tag拼接后重编码，作为测量范围变化的对照：
''')
    chunks.append(table(['设定','五项联合p','积极表达β／BH q','实际比较'], [[name,f(models[key]['joint_five_p']),f(coefs[key,'positive_expression']['coef'])+'／'+f(coefs[key,'positive_expression']['q_bh_five']),note]for key,name,note in [('baseline_body','原M3','正文线索，已有has_tag'),('body_tag_identity_controls','加9个tag身份控制','正文范围不变，调整常见标签差别'),('body_plus_tags_recoded','正文＋tag重编码','被测文字范围改变')]]))
    chunks.append('''拼接后，自律词命中从877增到3,342帖，问号从125增到354帖。统一的“坚持 Keep 的那些天”或带问号的活动标题可以机械地增加命中，不能等同为作者在正文中表现了更强自律或提出了问题。原始tag本身也不保证与正文主题一致，因此用作并行元数据和敏感性对照，不作为语义标签的标准答案。

9个常见tag是控制项，不是新的九类内容。没有进一步加入标签数量，是因为5,207帖只有一个tag，仅2帖有两个；在有无tag之外加数量，额外变化几乎只来自这两帖。

### 4.5 方向相近，不等于所有统计判断都稳定

原有的可见正文、去除最高1%点赞、作者等权和逐作者剔除分析中，积极表达的点估计方向较一致；100次逐作者剔除均为负，范围为[−0.100, −0.074]。PPML计数模型也给出负向点估计，但原始区间跨零。

![图6](figures/05_robustness.png)

图6　沿用原分析的三类词汇系数与原始95%作者聚类区间。原“仅可见正文”是5,548帖口径，不含本轮新增的5,414帖模型；不同尺度不能直接比较系数幅度。

![图7](figures/12_author_influence.png)

图7　原全样本设定中逐一剔除作者后的系数分布。这项检查处理单个作者影响，不处理语义效度或正文非随机缺失。

因此，本章保留的判断是：积极词的负向点估计在若干处理下相近，但整组证据与多重检验结论会随正文口径改变。数据提供的是受测量、样本和已记录控制条件限制的关联，不能据此建议通过减少积极表达来提高点赞。
''')
    chunks.append('''## 5. 五个用词成分不能直接当作完整内容分类

### 5.1 模型捕捉到反复出现的词群，覆盖边界仍然明显

原NMF最终使用2,470帖、86位作者。各成分的最高权重份额如下；为避免将模型结果误读为经过验证的五类内容，新图使用中性的成分编号与关键词。

![图8](figures/revised_component_shares.png)

图8　原冻结模型中最高权重成分的数值份额。每帖取一个最大权重用于描述，比例不是人工语义分类的精度或覆盖率；回归使用连续权重。

原名称分别为静蹲、骑行、运动训练交流、跑步打卡、生活感悟问候。其中有具体运动项目、上位内容域，也有表达功能，层级并不一致。52.7%与35.8%只能保留为原成分2和成分4的主导份额，不能直接认定为这些比例的帖子在语义上确实属于两个真实类别。

五类覆盖完整性的疑问有依据，但不能把2,944条未入模正文直接当作五类之外的帖子。未入模部分既有被词数门槛排除的明确运动或问候短句，也有饮食、装备等遗漏领域和信息不足的记录。

### 5.2 语义筛查发现具体错配，不能据此估计全样本错误率

本轮固定随机种子，对每个原主导成分抽12帖，共60帖，另抽20条未入模的非空正文；抽样不使用点赞。单轮AI辅助筛查发现，60条中40条明确支持原名称、6条只部分支持、12条不支持被分配的名称、2条信息不足。

例如，山河保护人的感悟、囤口罩提醒被分到静蹲；餐食配菜和食品品牌推荐被分到运动训练交流；绘画进展、出行和学期结束感受被分到跑步打卡。高归一化权重不能排除这种错配，因为它只给出几个成分之间的相对份额，不是类别正确的概率。[11]

各成分等额抽取使小成分被有意放大，不能用60条的未经加权比例推断全样本错误率。“不支持分配到的名称”也不等于不属于其余四类。本轮没有双人独立人工标注，不能声称已证明多数帖子都在五类之外。

### 5.3 事后给作者等权，不能消除拟合阶段的模板影响

原静蹲成分的85条主导文档中，78条来自一位作者；骑行成分的85条中，58条来自一位作者。沿用旧成分标签再计算作者平均份额，静蹲由3.4%降至0.6%，骑行由3.4%变为2.9%。这些数值说明产量与作者平均构成不同，但没有重新估计词群，因而不能证明分类已摆脱作者习惯。

2,478条拟合正文里有2,374种不同清理文本，104条为完全重复行。保持同一TF-IDF词表及IDF，删除完全重复正文后，原来独立的静蹲成分不再单列，相关词进入更宽的打卡成分；另出现包含三明治、咖啡、鸡蛋、早餐等词的成分。将拟合目标调整为作者等权时，词群结构也变化。作者等权诊断仍保留原始IDF，不是完整解决作者偏重的新文本管线。

K=3、4、6、7、8的对照与上述结果一起保留，未根据点赞、显著性或单纯拟合误差挑选新K。旧三次随机种子产生相近关键词，只证明同一设定容易复现；它不能代替语义、删重或作者构成稳健性。新出现的词群同样尚未完成语义验证。

### 5.4 连续权重可共存；二元多标记是另一种编码

原模型每帖有五项连续权重，合计为1，可以同时为正。带截距时省去成分2，是避免五项权重和为1带来的完全共线；省略不意味着语义互斥。回归系数表示其他权重和控制不变时，某一权重替代成分2权重的条件差异。

权重≥0.20的补充二元编码中，1,088帖有至少两个标记，成分2与成分4同时出现894帖。各标记比例合计145.9%，不能画成互斥饼图。五个标记同时进入回归时，设计矩阵171／171满秩，不能因存在五个变量就机械删除其中一个。
''')
    chunks.append(table(['权重阈值','至少两个标记的帖数','五标记联合p','最小BH q'], [[f"{x['threshold']:.2f}",f"{x['multi_topic_documents']:,}",f(x['joint_topic_p']),f(x['min_bh_q'])]for x in topics['threshold_results']]))
    chunks.append('''0.20取自1／5，只是便捷诊断；0.15与0.25用于观察阈值敏感性，并非语义金标准。三个阈值下都没有观测到“五项全零”组，二元系数不是相对于真实无主题组的差异。每个模型内五项BH也不覆盖整个阈值探索家族。给旧成分增加多标记没有修好其名称和覆盖问题，不能据此宣布已经获得新的真实主题回归。

### 5.5 成分回归仍是探索性关联，作者支持应在子样本内核算

原连续成分模型的联合p=0.166，四项BH校正后均未达到5%。成分4的原始p=0.035、BH q=0.107，不能据此给真实内容类型作稳定反馈排序。区间以冻结成分为给定输入，未纳入NMF估计与阈值选择的不确定性。

主题子样本有10位单帖作者，作者发帖数中位数为14；40位不超过10帖，合计181帖。不能用全样本只有2位单帖作者来替这个子模型保证支持充分。
''')
    chunks.append(table(['主题子样本内最低发帖数','保留帖数／作者数','原连续成分联合p'], [[x['minimum_topic_posts_per_author'],f"{x['n']:,}／{x['authors']}",f(x['joint_topic_p'])]for x in ta['models']]))
    chunks.append('''剔除少帖作者没有带来明确整组证据，但保留至少10条主题入模帖时只剩48位作者。它检查的是估计支持，并不修复成分的语义效度。也没有开展“积极表达×主题”的交互或分层检验，因此不能把积极词的关联反转归因于生活问候帖。

### 5.6 新的内容框架先统一分类轴，再做独立验证

候选一级内容域只回答“谈什么”，运动项目放到子层；“怎么表达”另列记录、经验建议、求助、问候或推荐等目的。帖文允许多标记，并分别保留其他可辨识内容、信息不足和无可用正文。
''')
    chunks.append(table(['候选内容域','拟覆盖的内容'], [[x['label'],x['includes']]for x in codebook['domains']]))
    chunks.append('''这份框架尚未用于生成全样本新标签，也不是新增一组已验证回归结果。下一步需两名同学独立标注覆盖作者、媒体、已入模和未入模帖的样本，讨论分歧后冻结边界，再用未参与修改的样本核查覆盖、各类一致性和“其他”占比。没有足够正文时保持未知，不仅凭原始tag补出主题；类数与边界也不由点赞或p值决定。
''')
    unchanged = old_section(old,'## 6. 鼓励与感谢有角色分工，回应覆盖另有差别','## 8. 对Keep社区基于数据的表达与互动的整体理解')
    chunks.append(unchanged)
    chunks.append('''## 8. 对表达与互动的整体理解

### 8.1 先区分观测到了什么，再解释点赞差异

原始均值、原M3和非空正文模型并非同一个比较。前者混合不同作者、媒体与发布情境；固定效应模型扣除已记录条件后给出关联；正文口径又决定哪些帖子的文字确实可以被测量。积极词的点估计在几个模型中为负，但五项联合检验p值从0.130、0.067变化到0.040，表明观察范围会影响结论。

四项字面特征没有明显共线性，少帖作者也没有主导原M3结果，但这两项检查不能替代内容测量。词典未命中不等于没有相关语义，tag存在不等于正文主题已知，NMF高权重也不等于名称判断正确。把这些边界分开，才能避免给局部结果加上过宽的“内容效果”解释。

作者和日期控制能改善已记录条件下的比较，但R²上升并不能分解真实的因果贡献，也不能证明“谁发、用什么媒体”比措辞更决定点赞。当前没有实际曝光、统一累计时长和完整媒体内容，不能据此安排增赞策略的优先级。

### 8.2 评论的活跃，主要体现为围绕帖主的留言与回应

沿用的评论结果显示，其他用户更常出现赞许词，帖主更常出现感谢词；多数明确回复由帖主发给其他用户。直接配对把部分赞许留言与感谢回复联系起来，但它仍是已采集评论的描述，不能解释为鼓励留言造成更多回复。

评论总量、参与范围和回应覆盖可以彼此分离。接收事件量相近的U1528与U0948分别来自53人和1,091人；鼓励组在已回复内容中更常出现感谢，但根评论的观测直接回复比例更低。研究互动时，应同时看谁参与、向谁发言、往来了多少次及多少留言得到承接。

### 8.3 文字关联与评论结构尚未组成经验证的机制链条

现有分析分别描述字面线索与点赞关联、评论语言和回复网络。它没有检验特定文字引起何种回复、再由回复带来更多点赞，也没有验证积极词反转集中于某个内容主题。接下来的解释应从完整帖文、实际曝光与直接承接记录出发，而不是用一个负向成分系数替代机制检验。

## 9. 建议、局限与后续研究

### 9.1 当前能够采用的判断

数据足以描述样本内的长尾点赞、媒体组差别、字面规则的覆盖，以及已采集评论的角色和连接结构。字面关联则对正文可观察范围敏感，五个用词成分也尚不能代表完整真实内容分类。因此，本报告不推出固定关键词、减少积极措辞或选择某个主题能够稳定提高点赞。

对于内容创作者，当前结果改变的是比较方式：不能只看高赞帖常用什么词，就把作者、媒体与观察时长的差别当作措辞收益。对于研究者和平台运营者，应先提高内容与曝光记录的可比较性，再评估表达提示或互动干预。

### 9.2 可参考的工作方向

| 对象 | 下一步可做什么 | 证据依据与条件 |
| --- | --- | --- |
| 研究小组 | 冻结正文可观察范围，分开字面、原始tag与人工主题 | 本轮口径对照改变统计判断；不同表示不能混用 |
| 研究小组 | 对候选内容框架开展双人独立标注和留出验证 | 语义筛查发现错配，删重改变成分；当前尚无金标准 |
| 平台分析者 | 统一记录发帖后7天曝光、反馈和不同参与者数 | 现有累计点赞时长未知，事件数与参与人数含义不同 |
| 平台运营者 | 将内容或回复提示作为可检验干预，而非既定有效策略 | 观察性结果没有识别因果；实验方案见9.4 |

### 9.3 仍限制结论的条件

样本来自课程提供的100位作者，抽样机制未知，不能代表Keep全平台。正文缺失和主题入模都带来选择，缺失标记与仅正文分析均不能恢复未记录内容；图片文字和视频语音尚未观察。

三类词汇加问号不是经过验证的心理量表或语义分类器，原始tag也不是人工真值。五个NMF成分对模板和作者产量敏感，单轮AI筛查不能估计全样本分类精度。候选框架尚未完成独立人工验证。

正文口径、作者门槛、tag和阈值对照是在原结果后开展的探索；各模型内的BH不覆盖全部模型探索。固定效应与作者聚类不能解决实际曝光、粉丝随时间变化、媒体内容和累计时长的缺失。点赞不等于训练效果、自律程度或留存。

评论与网络结果来自已采集互动，不能把未采到回复视为实际没有回复。第6—7章与相关旧图本轮未全面重审，帖子模型的复核不代表这些部分已获得新的完整验证。

### 9.4 后续研究设计

先完成候选内容域与表达目的的独立标注和留出核查，再以冻结规则开展新的主题比较。需要公布各类覆盖、混淆与未知状态，保留饮食、健康、装备、平台和日常等真实出现的内容，不为提高覆盖而无限扩大生活类。

若获得平台实验权限，可按作者随机分组比较运动记录提示、面向伙伴的问题提示与常规发布，并统一7天观察窗口，记录实际曝光、独立互动用户数、后续发布和回复质量。样本量应在试点后按方差及最小有意义差异估算。这是尚未实施的后续设计，不是当前关联结果证明有效的做法。

### 9.5 结论

本研究保留可复算的字面与互动描述，同时收窄了内容结论。四项字面规则的共线性不突出、少帖作者不是原结果的主要来源；但文字关联随正文可观察范围变化，五个用词成分也尚不足以作为完整真实内容分类。评论部分则显示，在已采集记录中，回话主要围绕帖主，互动次数与参与人数反映不同面貌。现有证据适合提出更明确的测量和验证问题，不适合给关键词或主题作稳定的增赞排序。
''')
    chunks.append('''## 附录A　模型口径与复现证据

### A.1 原全样本模型与非空正文模型的完整五项结果

以下原始区间是各单项95%作者聚类区间，BH q在每个模型内对五项校正。修订正文范围是5,414帖、95位作者；原M3是9,532帖、100位作者，作为历史基线保留。
''')
    rows=[]
    for model,label in [(cm[0],'原M3'),(cm[2],'非空正文')]:
        for x in model['terms']:
            rows.append([label,LABELS[x['term']],f(x['coef']),f(x['se_cluster']),ci(x),f(x['p']),f(x['q_bh'])])
    chunks.append(table(['模型','字面特征','β','聚类SE','原始95%区间','原始p','BH q'],rows))
    chunks.append('''### A.2 原有敏感性分析的保留范围

原有可见正文、最高1%剔除、作者等权与PPML结果仍在冻结结果表中，五项变量和控制项定义未改。原“仅可见正文”为5,548帖，不能与新增5,414帖混称。
''')
    app2=old_section(old,'### A.2 积极表达词的敏感性结果','### A.3 主题描述与主题回归')
    chunks.append(app2[app2.index('| 补充设定'):])
    chunks.append('''### A.3 原连续成分回归及二元多标记口径

原TF-IDF词表含1,046项，NMF拟合2,478帖，8帖近零贡献被排除，最终2,470帖。连续权重在每帖合计为1，省略成分2作参照。下表名称仅用于对照历史结果，不表示类别已验证。
''')
    reg=pd.read_csv(ROOT/'results/extension_topic_regression.csv').set_index('topic')
    rows=[]
    for k in range(5):
        if k==2: rows.append([f'成分{k}（旧称{topic_names[k]}）',topics['original_dominant_counts'][k],'参照','不适用','不适用','不适用'])
        else:
            x=reg.loc[k].to_dict()
            rows.append([f'成分{k}（旧称{topic_names[k]}）',topics['original_dominant_counts'][k],f(x['coef']),ci(x),f(x['p']),f(x['q_bh'])])
    chunks.append(table(['冻结成分','主导帖数','连续权重β','原始95%区间','原始p','BH q'],rows))
    chunks.append('''连续系数按1单位权重变化计；解释10个百分点替代时需乘0.1。二元标记则定义为权重达到阈值，表示固定其余四个标记和控制后的条件差异；五项同时进入、没有单一参照类。三个阈值均没有真实全零组。所有区间未纳入主题估计、语义归类与阈值选择的不确定性。

### A.4 评论配对、网络与时间口径
''')
    a4=old_section(old,'### A.4 评论配对、网络与时间口径','### A.5 媒体分组的完整描述数值')
    chunks.append(a4[a4.index('\n')+1:])
    a5=old_section(old,'### A.5 媒体分组的完整描述数值','## 附录B')
    chunks.append(a5)
    chunks.append('''### A.6 本轮新增证据的可复现位置

| 分析内容 | 汇总证据 |
| --- | --- |
| 分类理由、共线性与作者内变化 | results/teacher_feedback_dictionary.json及相关、VIF表 |
| 三种正文口径完整结果 | results/teacher_complete_body_model.json |
| 主样本作者门槛、tag身份和拼接 | results/teacher_author_tags_summary.json及系数表 |
| 连续权重与阈值多标记 | results/teacher_topic_summary.json及模型、系数表 |
| 主题子样本作者支持 | results/teacher_topic_author_support.json |
| 覆盖和单轮AI语义筛查 | results/teacher_label_coverage_summary.json |
| K、删重和作者拟合权重 | results/teacher_topic_structure_summary.json及关键词表 |
| 未验证的候选编码规则 | code/revised_content_codebook.json |

原始帖ID与正文筛查表仅保留本地，公开材料使用聚合统计和无原始帖ID的判定。完整复现入口沿用本项目Makefile；本报告生成脚本只读已有证据，不重算分析。正文、图表和附录均使用相同冻结结果。

## 附录B　成员分工

沿用当前正式汇报中已确认的实际分工：陈霖负责问题提出、数据清洗与RQ1；董依心负责RQ2与RQ3。本版不沿用早期建议框架中的空白成员或占位学号。
''')
    refs=old[old.index('## 参考文献'):]
    refs=re.sub(r'(?<=\w)<br>\s*(?=\w)',' ',refs).replace('<br>\n','')
    for wrapped,word in [('Inter-net','Internet'),('Process-ing','Processing'),('Re-sources','Resources'),('Cam-bridge','Cambridge')]:
        refs=refs.replace(wrapped,word)
    refs += '''\n\n[10] statsmodels. Variance inflation factor. 官方方法说明：[VIF文档](https://www.statsmodels.org/v0.14.6/generated/statsmodels.stats.outliers_influence.variance_inflation_factor.html)。\n\n[11] scikit-learn. NMF. 官方实现说明：[NMF文档](https://scikit-learn.org/1.5/modules/generated/sklearn.decomposition.NMF.html)。本项目的贡献归一化、样本筛选和语义诊断另见代码与冻结结果；实现文档本身不能验证数据的类别名称。\n'''
    chunks.append(refs)
    return '\n\n'.join(chunks).strip()+'\n'


PDF_CSS = '''body{font-family:sans-serif;font-size:10.5pt;line-height:1.45;color:#20332f;}h1{font-size:23pt;color:#173f36;}h2{font-size:17pt;color:#245c4f;margin-top:16pt;}h3{font-size:12.5pt;color:#245c4f;margin-top:13pt;}p{margin:7pt 0;}table{border-collapse:collapse;width:100%;font-size:8.5pt;margin:9pt 0;}th,td{padding:4pt;border:0;vertical-align:top;}th{font-weight:bold;color:#245c4f;}img{width:100%;}figure{margin:9pt 0;}figcaption{font-size:9pt;color:#536b61;}a{color:#245c4f;}'''


def export(md):
    (ROOT/f'{STEM}.md').write_text(md)
    office.STEM=STEM; office.VALIDATION=QA
    office.CSS=PDF_CSS+' .document{max-width:940px;margin:25px auto;padding:35px;background:white;} .toc a{display:block;} .toc .level-2{margin-left:15px;} .chapter{margin-top:30px;}'
    _,soup,headings=office.read_manuscript()
    final=office.write_html(soup,headings)
    office.write_docx(final,headings)
    # Separate stories make each chapter begin on a page without relying on
    # browser CSS pagination. Tables have no painted backgrounds or borders.
    sections=final.select('main > section')
    toc=final.select_one('nav.toc')
    pieces=[sections[0],toc]+sections[1:]
    target=ROOT/f'{STEM}.pdf'
    paper=fitz.paper_rect('a4'); area=paper+(44,42,-44,-44)
    def render_once(lookup=None):
        writer=fitz.DocumentWriter(str(target))
        starts=[]; page_count=0
        for node in pieces:
            if node is None: continue
            title=node.find(['h1','h2']).get_text() if node.find(['h1','h2']) else '目录'
            if node.name=='nav':
                rows=[]
                for section in sections[1:]:
                    heading=section.find('h2').get_text()
                    number=str(lookup[heading]) if lookup else '—'
                    rows.append('<tr><td>'+html.escape(heading)+'</td><td>'+number+'</td></tr>')
                content='<h2>目录</h2><table><tr><th>章节</th><th>页码</th></tr>'+''.join(rows)+'</table><p>小节标题见各章正文。PDF书签可直接跳转章节。</p>'
            else:
                content=str(node)
            section_css=PDF_CSS
            if title.startswith('9.'):
                # Keep the short concluding chapter together without changing
                # body sizing or pagination in any earlier chapter.
                section_css+=' body{line-height:1.37;}p{margin:6pt 0;}h3{margin-top:11pt;}'
            markup='<html><head><style>'+section_css+'</style></head><body>'+content+'</body></html>'
            story=fitz.Story(html=markup,archive=fitz.Archive(str(ROOT)))
            starts.append({'title':title,'page':page_count+1})
            more=1
            while more:
                device=writer.begin_page(paper)
                more,_=story.place(area); story.draw(device); writer.end_page(); page_count+=1
                assert page_count<80
        writer.close()
        return starts
    first_starts=render_once()
    starts=render_once({x['title']:x['page']for x in first_starts})
    assert starts==first_starts,'Adding TOC page numbers unexpectedly changed pagination'
    pdf=fitz.open(target)
    for i,page in enumerate(pdf,1):
        page.insert_text((paper.width/2-8,paper.height-24),str(i),fontsize=9,color=(.35,.4,.37))
    pdf.set_toc([[1,x['title'],x['page']]for x in starts if x['title']!='目录'])
    temp=QA/'numbered_report.pdf'; pdf.save(temp,garbage=4,deflate=True); pdf.close(); os.replace(temp,target)
    return starts


def validate(md, starts, protected):
    from docx import Document
    pdf=fitz.open(ROOT/f'{STEM}.pdf')
    compact=re.sub(r'\s+','',''.join(p.get_text()for p in pdf))
    doc=Document(ROOT/f'{STEM}.docx')
    docx_text=''.join(p.text for p in doc.paragraphs)+''.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    required=['71.9%','2,944','1.11—1.75','0.130','0.067','0.040','0.048','1,088','145.9%','0.102','0.173','10位单帖作者','单轮AI','候选','89.0%','1,091','未开展','陈霖','董依心']
    # The manuscript says "没有开展" rather than "未开展"; both mean no test.
    required.remove('未开展')
    missing=[x for x in required if re.sub(r'\s+','',x)not in compact]
    assert not missing, missing
    assert all(x in docx_text for x in ['0.040','0.048','1,088','董依心','候选'])
    assert not re.search(r'【待填写|\bTODO\b|lorem ipsum',md)
    assert '反转可能集中出现在日常问候类帖子中'not in md
    assert '点赞差异主要对应作者和内容形式'not in md
    images=re.findall(r'!\[[^]]*\]\(([^)]+)\)',md)
    assert len(images)==14 and all((ROOT/x).is_file()for x in images),images
    original=SOURCE.read_text()
    unchanged=old_section(original,'## 6. 鼓励与感谢有角色分工，回应覆盖另有差别','## 8. 对Keep社区基于数据的表达与互动的整体理解')
    assert unchanged in md,'Comment and network chapters must remain verbatim'
    over=[]
    for old_image in QA.glob('page-[0-9][0-9].png'):
        old_image.unlink()
    for n,page in enumerate(pdf,1):
        assert page.get_text().strip(),n
        for b in page.get_text('dict')['blocks']:
            for line in b.get('lines',[]):
                for span in line.get('spans',[]):
                    x0,y0,x1,y1=span['bbox']
                    if x0<20 or y0<10 or x1>page.rect.width-20 or y1>page.rect.height-10:
                        over.append({'page':n,'text':span['text'],'bbox':span['bbox']})
        page.get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).save(QA/f'page-{n:02d}.png')
    assert not over,over[:5]
    for batch,start in enumerate(range(0,len(pdf),9),1):
        count=min(9,len(pdf)-start)
        sheet=Image.new('RGB',(1500,((count+2)//3)*720),'#e3e9e5');draw=ImageDraw.Draw(sheet)
        for j in range(count):
            im=Image.open(QA/f'page-{start+j+1:02d}.png');im.thumbnail((480,680))
            x=(j%3)*500+8;y=(j//3)*720+26;sheet.paste(im,(x,y));draw.text((x,y-19),str(start+j+1),fill='#173f36')
        sheet.save(QA/f'contact-{batch}.png')
    assert all(hashlib.sha256(p.read_bytes()).hexdigest()==v for p,v in protected.items())
    manifest={'source':str(SOURCE.relative_to(ROOT)),'pdf_pages':len(pdf),'figures':len(images),'docx_tables':len(doc.tables),'docx_inline_images':len(doc.inline_shapes),'required_claims_present':True,'comments_and_network_verbatim':True,'text_outside_safe_page':over,'sources_unchanged':True,'chapter_starts':starts,'visual_review':'pending','new_report_sha256':{ext:hashlib.sha256((ROOT/f'{STEM}.{ext}').read_bytes()).hexdigest()for ext in ['md','html','docx','pdf']},'protected_inputs':{str(p.relative_to(ROOT)):v for p,v in protected.items()}}
    (QA/'report_checks.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in manifest.items()if k!='protected_inputs'},ensure_ascii=False,indent=2))


def main():
    QA.mkdir(parents=True,exist_ok=True)
    protected_paths=[SOURCE,ROOT/'方法复核与修订结论.md']+list((ROOT/'results').glob('teacher_*.json'))+list((ROOT/'figures').glob('[0-9][0-9]_*.png'))
    protected={p:hashlib.sha256(p.read_bytes()).hexdigest()for p in protected_paths}
    md=manuscript(); starts=export(md); validate(md,starts,protected)


if __name__=='__main__':
    main()
