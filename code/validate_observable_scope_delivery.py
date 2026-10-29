"""Cross-check the delivered scope against unchanged numerical evidence."""
from pathlib import Path
import hashlib
import json
import re

import pymupdf as fitz
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
QA=ROOT/'validation/observable_scope'
REPORT='分析报告_字面特征与互动结构版'
SUMMARY='本轮修订说明_研究范围与结论边界'


def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def compact(s):return re.sub(r'\s+','',s)


def main():
    baseline=json.loads((QA/'protected_inputs.json').read_text())
    assert all(digest(ROOT/p)==h for p,h in baseline.items()),'A frozen input changed.'
    scope=json.loads((ROOT/'code/observable_feature_scope.json').read_text())
    state=json.loads((ROOT/'results/current_study_scope.json').read_text())
    assert state['scope_definition_sha256']==digest(ROOT/'code/observable_feature_scope.json')
    assert scope['human_semantic_review']['completed_human_coders']==0
    assert not scope['human_semantic_review']['current_delivery_requires_human_annotation']
    assert state['human_semantic_review']['status']=='not_conducted_in_current_study'
    md=(ROOT/(REPORT+'.md')).read_text()
    headings=[m for m in re.finditer(r'^##\s+(.+)$',md,re.M)]
    appendix=next((m for m in headings if '附录'in m.group(1)),None)
    assert appendix,'Exploratory material must be separated into appendices.'
    main_text=md[:appendix.start()]
    assert all(x in main_text for x in ['坚持等词项','大家等词项','开心等词项','问号'])
    for x in ['自律／目标词','自律/目标词','社交／互助词','社交称呼/互助词','积极表达词命中']:
        assert x not in main_text,('old focal category name remains',x)
    assert any(x in main_text for x in ['未开展人工语义','未进行人工语义','未开展独立人工','未进行独立人工'])
    assert all(x in main_text for x in ['0.130','0.067','0.040'])
    assert '事前'in main_text and '每个模型'in main_text
    assert '探索'in md[appendix.start():]and 'NMF'in md[appendix.start():]
    assert '24'in md[appendix.start():]and '18'in md[appendix.start():]
    assert not any(x in main_text for x in ['下一步需两名同学','先完成候选内容域','先完成双人独立标注'])
    output={}
    for stem in [REPORT]:
        pdf=fitz.open(ROOT/(stem+'.pdf'))
        text=''.join(p.get_text()for p in pdf)
        assert '\ufffd'not in text
        assert all(len(p.get_text().strip())>15 for p in pdf)
        assert '未开展'in text or '未进行'in text
        word=Document(ROOT/(stem+'.docx'))
        if stem==REPORT:
            assert len(word.tables)>=5
            assert all(v in compact(text)for v in ['0.130','0.067','0.040','150','41','37'])
            images=re.findall(r'!\[[^]]*\]\(([^)]+)\)',md)
            assert images and all((ROOT/p).is_file()for p in images)
        else:images=[]
        output[stem]={'pages':len(pdf),'editable_word_tables':len(word.tables),'figures':len(images),'all_pages_have_text':True,'pdf_sha256':digest(ROOT/(stem+'.pdf'))}
        pdf.close()
    assert '当前正式PDF报告'in(ROOT/'README.md').read_text()
    assert '本轮未开展人工语义复核'in(ROOT/'人工编码开始这里.md').read_text()
    result={'status':'pass','frozen_input_count':len(baseline),'frozen_inputs_unchanged':True,'model_results_reused_without_reselection':True,
            'human_semantic_review':'not_conducted_in_current_study','current_delivery_requires_human_annotation':False,
            'delivery_format':'pdf_report','classification_material_in_exploratory_appendices':True,'report':output,
            'limitation':'本检查核对输入、数值呈现、文稿结构和PDF；不代表人工语义效度验证。'}
    QA.mkdir(exist_ok=True)
    (QA/'delivery_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
